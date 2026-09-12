"""
"How these figures are calculated", as a document.

WHAT MAKES IT CONVINCING, held as tests.

Not prose about rigour. A challenged figure needs three things: the RULE it
was computed by, the VALUES that went into the rule, and the rule with those
values substituted so the arithmetic can be followed to the number on the
screen. Every check below is about one of those three, or about the honesty of
the cases where one of them is missing.
"""

from __future__ import annotations

import io
import pathlib

import pytest

from netgravity.orchestrator.agents.llm_gateway import LLMResponse
from netgravity.orchestrator.metrics.warehouse_deep_dive import (
    WarehouseDeepDiveReport,
    WarehouseHealthKPI,
)
from netgravity.reporting.derivation import build_derivation_docx
from netgravity.reporting.kpi_method import build_kpi_method_report
from netgravity.reporting.narration import allowed_figures, narrate

REPO_ROOT = pathlib.Path(__file__).resolve().parents[3]


def _kpis():
    return {
        "business_network_cost": {"value": 150627.70, "status": "VALID",
                                  "display_value": "INR 150,627.70"},
        "transport_cost": {"value": 45889.70, "status": "VALID"},
        "facility_cost": {"value": 95000.0, "status": "VALID"},
        "handling_cost": {"value": 8030.0, "status": "VALID"},
        "inventory_cost": {"value": 1708.0, "status": "VALID"},
        "demand_fill_rate": {"value": 1.0, "status": "VALID",
                             "display_value": "100.0%"},
        "served_demand": {"value": 21900.0, "status": "VALID"},
        "total_demand": {"value": 21900.0, "status": "VALID"},
    }


def _report():
    row = WarehouseHealthKPI(
        facility_id="DC_CENTRAL", facility_name="Central Distribution Centre",
        role="DC", is_open=True, rated_capacity_per_period=5000,
        peak_throughput_units=2700, peak_utilization_pct=54.0)
    return WarehouseDeepDiveReport(health_kpis=[row])


def _built(**overrides):
    kwargs = dict(kpis=_kpis(), report=_report(), threshold=90.0,
                  currency="INR",
                  scope={"project_name": "Demo", "lens_label": "Every facility"})
    kwargs.update(overrides)
    return build_kpi_method_report(**kwargs)


class TestTheArithmeticIsShown:

    def test_every_calculation_carries_a_formula(self):
        equations = [e for s in _built().steps for e in s.equations]
        assert equations
        for equation in equations:
            assert equation.formula, equation.name

    def test_the_symbols_are_defined(self):
        """A formula with undefined symbols is decoration."""
        for step in _built().steps:
            for equation in step.equations:
                if "=" not in equation.formula:
                    continue
                if not equation.variables:
                    # The banding rule defines its own terms in the formula.
                    assert "if" in equation.formula.lower(), equation.name
                    continue
                for variable in equation.variables:
                    assert variable.symbol and variable.meaning, equation.name

    def test_the_cost_substitution_sums_to_the_stated_total(self):
        """
        A document whose worked example disagrees with its own results table
        is worse than one with no example.
        """
        step = _built().steps[0]
        worked = step.equations[0].substituted
        assert "45,890" in worked and "95,000" in worked
        assert "8,030" in worked and "1,708" in worked
        assert "150,628" in worked
        # And the components genuinely add up to it.
        assert round(45889.70 + 95000.0 + 8030.0 + 1708.0, 2) == 150627.70

    def test_the_utilisation_substitution_follows_its_own_formula(self):
        step = next(s for s in _built().steps if "working" in s.title)
        equation = step.equations[0]
        assert equation.formula == "U_peak = (T_peak / C_rated) x 100"
        assert "(2,700 / 5,000) x 100 = 54.0%" in equation.substituted
        assert round((2700 / 5000) * 100, 1) == 54.0

    def test_the_worked_example_names_the_site_it_came_from(self):
        step = next(s for s in _built().steps if "working" in s.title)
        assert "Central Distribution Centre" in step.equations[0].note


class TestAbsenceIsSaidNotSubstituted:

    def test_a_network_with_no_readings_still_states_the_rule(self):
        """
        "We cannot compute this, and here is the rule we would have used" is
        an answer. A formula quietly dropped leaves the reader assuming the
        metric does not exist.
        """
        built = _built(kpis={}, report=WarehouseDeepDiveReport())
        equations = [e for s in built.steps for e in s.equations]
        assert equations
        assert any("U_peak" in e.formula for e in equations)

    def test_a_missing_value_is_never_shown_as_zero(self):
        built = _built(kpis={}, report=WarehouseDeepDiveReport())
        blob = " ".join(
            f"{e.substituted} {' '.join(v.value for v in e.variables)}"
            for s in built.steps for e in s.equations)
        assert "Not available" in blob
        assert "= 0.0%" not in blob
        assert "= 0 units" not in blob

    def test_an_invalid_kpi_is_absent_rather_than_read(self):
        """The layer marks a figure VALID or it means nothing."""
        kpis = _kpis()
        kpis["business_network_cost"] = {"value": 999.0, "status": "NOT_COMPUTED"}
        built = _built(kpis=kpis)
        assert "999" not in built.steps[0].equations[0].substituted


class TestThePeriodConversionIsStated:
    """
    The one nobody expects, and the one that produced a real defect: a year's
    fixed cost charged as a month's is twelve times the truth, and made
    opening one site look like it doubled the cost of the network.
    """

    def test_the_conversion_has_its_own_step(self):
        step = next((s for s in _built().steps if "fixed cost" in s.title.lower()),
                    None)
        assert step is not None
        assert "C_fixed = (F_annual / 12) x N_periods" in step.equations[0].formula

    def test_it_says_how_the_period_is_decided(self):
        step = next(s for s in _built().steps if "fixed cost" in s.title.lower())
        note = step.equations[0].note
        assert "heading" in note
        assert "twelve" in note


class TestTheDocumentRenders:

    def test_it_builds_a_docx_with_the_equations_in_it(self):
        blob = build_derivation_docx(_built())
        assert blob[:2] == b"PK"          # a .docx is a zip
        assert len(blob) > 10_000

        import docx
        document = docx.Document(io.BytesIO(blob))
        text = "\n".join(p.text for p in document.paragraphs)
        assert "U_peak = (T_peak / C_rated) x 100" in text
        assert "C_network =" in text
        assert "F = (D_served / D_total) x 100" in text
        # One table per equation that defines symbols, plus the figure tables.
        assert len(document.tables) >= 6

    def test_it_states_what_it_does_not_establish(self):
        built = _built()
        assert built.limitations
        joined = " ".join(built.limitations)
        assert "shortage penalty" in joined
        assert "do not sum" in joined


class TestTheProseIsGroundedInTheArithmetic:

    class _Gateway:
        available = True

        def __init__(self, output):
            self._output = output

        def generate(self, prompt, purpose=""):
            self.prompt = prompt
            return LLMResponse(output=self._output)

    def test_a_figure_from_an_equation_may_be_cited(self):
        """
        `allowed_figures` used to read only the figures TABLES, so any
        sentence walking the reader through a calculation was struck out —
        which is the one thing a methodology document is written to do.
        """
        allowed = allowed_figures(_built())
        assert 5000.0 in allowed          # a variable's value
        assert 54.0 in allowed            # a substitution's result
        assert 2700.0 in allowed          # a substitution's operand

    def test_an_invented_figure_is_struck_out(self):
        gateway = self._Gateway(
            "Central Distribution Centre runs at 54.0% of its 5,000 unit "
            "capacity. The network cost is 999,999 which nobody computed.")
        narration = narrate(_built(), gateway, purpose="test")
        kept = " ".join(narration.paragraphs)
        assert "54.0%" in kept
        assert "999,999" not in kept
        assert narration.dropped

    def test_nothing_is_claimed_when_no_gateway_is_configured(self):
        """
        The default in every environment without a credential. A document with
        no narration is a COMPLETE document; what it must not do is imply a
        model wrote the template's words.
        """
        narration = narrate(_built(), None, purpose="test")
        assert narration.paragraphs == ()
        assert narration.source == "unavailable"
        assert narration.note


class TestTheRouteIsWired:

    def _api(self) -> str:
        return (REPO_ROOT / "app" / "backend" / "api"
                / "kpis.py").read_text(encoding="utf-8")

    def test_the_document_route_exists_and_narrates(self):
        api = self._api()
        assert '@bp.route("/method.docx", methods=["POST"])' in api
        block = api[api.index("def export_kpi_method():"):]
        block = block[:block.index("\n    @bp.route")]
        assert "build_kpi_method_report" in block
        assert "narrate(derivation, _gateway()" in block

    def test_it_uses_the_shared_gateway_not_a_second_one(self):
        """
        The budget is cumulative across every holder of the token — 100
        requests a day for the whole product.
        """
        api = self._api()
        block = api[api.index("def _gateway()"):]
        block = block[:block.index("\n    @bp.route")]
        assert 'services or {}).get("reasoning_agent")' in block
        # The CODE, not the docstring explaining what the code avoids —
        # which names `LLMGateway()` in order to say it is not
        # constructed here. This is the fourth time in this codebase a
        # check has matched its own explanation.
        import re
        code = re.sub(r'\"\"\".*?\"\"\"', '', block, flags=re.S)
        assert "LLMGateway()" not in code

    def test_the_button_is_on_the_screen(self):
        html = (REPO_ROOT / "app" / "frontend"
                / "index.html").read_text(encoding="utf-8")
        panel = html[html.index('id="tab-facility-dashboard"'):]
        panel = panel[:panel.index("</section>")]
        assert 'id="btn-export-method"' in panel
        assert "How it is calculated" in panel
