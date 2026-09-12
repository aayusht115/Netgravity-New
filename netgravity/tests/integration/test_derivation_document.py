"""
The derivation document, and the figures that reach a reader.

Two things are held here.

  * **A finding can leave the application.** The deep dive shows the
    conclusion, the figures it cites and the role each played — the right
    amount for a screen and the wrong amount for the conversation that
    follows it. The first question asked of a capacity finding in a steering
    committee is which figures it rests on and what the model could not see,
    and the answer has to survive being forwarded to somebody who will never
    open this application.

  * **Nothing in it is gibberish.** Every caption on every screen used to be
    the storage key in Title Case — "N Facilities Open", "Avg Utilization
    Pct", "Pct Demand In Sla" — and a demand fill rate was printed as the
    ratio "1.000". Those strings are the caption under the headline figure on
    an Overview tile, the label on every Insights row, and the first column
    of the table in the document that leaves the building.

The writer itself is `netgravity.reporting`, which knows nothing about
insights: the demand forecast is asked the same question and will build the
same shape, so the tests here separate the two.
"""

from __future__ import annotations

import io

import pytest

from netgravity.orchestrator.reasoning.evidence import (
    _display, build_evidence_pack, metric_label, twin_reasoning_payload,
    with_policy_thresholds,
)
from netgravity.orchestrator.registry import build_orchestrator
from netgravity.orchestrator.schemas.reasoning import ReasoningScope
from netgravity.orchestrator.schemas.requests import (
    Actor, ActorRole, Intent, OrchestratorRequest)
from netgravity.reporting import (
    DerivationReport, DerivationStep, Figure, build_derivation_docx)
from netgravity.tests.fixtures.case16_synthetic import build_case16_network

import app.backend.api.insights as insights_api


@pytest.fixture(scope="module")
def briefing_bundle():
    """A real solved network, its briefing, and the serialised findings."""
    orc = build_orchestrator()
    snapshot = orc.snapshots.register(build_case16_network(), label="docs")
    orc.run_sync(OrchestratorRequest(
        input="baseline", explicit_intent=Intent.NETWORK_STATE_QUERY,
        actor=Actor(actor_id="u", role=ActorRole.PLANNER),
        network_snapshot_id=snapshot.snapshot_id, disable_llm=True))
    state = orc.twin.materialize(
        orc.twin.list_states(snapshot.snapshot_id)[-1].state_id)

    payload = with_policy_thresholds(twin_reasoning_payload(
        state, scope=ReasoningScope.NETWORK, entity_id=None, comparison=None))
    unavailable = {i.field: {"status": i.status.value, "reason": i.reason}
                   for i in state.unavailable}
    pack = build_evidence_pack(
        payload, scope=ReasoningScope.NETWORK, entity_id=None,
        user_question="", unavailable=unavailable,
        provenance={"state_id": state.state_id})
    result = orc.services["reasoning_agent"].reason(
        payload, unavailable_evidence=unavailable,
        provenance={"state_id": state.state_id}, allow_llm=False,
        scope=ReasoningScope.NETWORK, entity_id=None, user_question="")

    records = [
        insights_api._serialise_insight(item, i, scope="NETWORK",
                                        entity_id=None, pack=pack)
        for i, item in enumerate(result.briefing.kpi_insights)
    ]

    # THE SHAPE THE ENDPOINT SERVES, because that is what the document is now
    # built from. `_derivation_for` used to take the live briefing, result and
    # twin state; it takes the serialised payload instead, so the document and
    # the screen read one object and cannot disagree about what findings exist
    # (see the route's own comment for the 404 this fixed).
    analysis = {
        "state_id": state.state_id,
        "insights": records,
        "key_drivers": list(result.briefing.key_drivers or []),
        "limitation": result.briefing.limitation or "",
        "evidence_completeness": getattr(
            result.briefing.evidence_completeness, "value",
            str(result.briefing.evidence_completeness)),
        "grounding": {
            "status": result.grounding_status,
            "warnings": list(result.validation_warnings),
            "source": result.source,
        },
    }
    return result, records, state, analysis


def _docx_text(data: bytes) -> str:
    """Everything a reader would see in the file, as one string."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


class TestTheWriterKnowsNothingAboutInsights:
    def test_it_renders_a_report_built_by_hand(self):
        """
        The point of the split. A demand forecast is asked the same question —
        which series, which method, what history — and will build this shape
        without the insights endpoint being involved.
        """
        report = DerivationReport(
            kind="Demand forecast",
            subject="Toronto Metro · P001",
            conclusion="Demand grows 4% over the next two quarters",
            summary="The recorded history supports a mild upward trend.",
            method="Fitted on 24 recorded months.",
            steps=[DerivationStep(
                title="What was measured",
                detail="The history the forecaster was given.",
                figures=(Figure("Observed months", "24", "Measured", "forecast_engine"),),
            )],
            recommended_action="Test the footprint against the projected demand.",
            limitations=["Two markets state no history."],
            provenance="Source: forecasting engine.",
        )
        text = _docx_text(build_derivation_docx(report))
        assert "Demand grows 4% over the next two quarters" in text
        assert "Observed months" in text and "24" in text
        assert "Two markets state no history." in text
        assert report.filename().endswith(".docx")
        assert "Demand-forecast" in report.filename()

    def test_a_section_with_nothing_in_it_is_omitted(self):
        """
        An "Assumptions" heading over a blank suggests the model made none,
        which is a claim — and a different one from "none were recorded".
        """
        text = _docx_text(build_derivation_docx(DerivationReport(
            subject="s", conclusion="c")))
        assert "What the model was given" not in text
        assert "What this does not establish" not in text
        assert "Recommended action" not in text


class TestAFindingBecomesADocument:
    def test_every_figure_matches_the_screen_exactly(self, briefing_bundle):
        """
        The document and the deep dive print the identical `display_value`
        string, so the file and the screen it came from cannot disagree about
        a number. A document that reformatted — or worse, recomputed — one
        would be a second engine whose output carries a letterhead and gets
        forwarded.
        """
        _result, records, _state, analysis = briefing_bundle
        record = next(r for r in records if r["theme"] == "Capacity")
        report = insights_api._derivation_for(record, analysis)
        text = _docx_text(build_derivation_docx(report))

        cited = [e for e in record["evidence"]
                 if e["display_value"] and e["display_value"] != "Not available"]
        assert cited, record
        for row in cited:
            assert row["display_value"] in text, row
            assert row["label"] in text, row
        assert record["headline"] in text
        assert record["recommended_action"] in text

    def test_it_carries_the_records_the_screen_could_only_draw(self, briefing_bundle):
        """
        The entities a finding was computed over are a chart on screen and a
        list in the file — which site, at what figure. That is the part a
        reader is asked for first and the part a bar chart cannot be pasted
        into a deck as.
        """
        _result, records, _state, analysis = briefing_bundle
        record = next(r for r in records if r["entities"])
        report = insights_api._derivation_for(record, analysis)
        text = _docx_text(build_derivation_docx(report))
        for entity in record["entities"][:3]:
            assert str(entity["label"]) in text, entity

    def test_it_says_whether_the_figures_were_checked(self, briefing_bundle):
        """
        A reader acting on prose is entitled to know whether its numbers were
        verified against the deterministic results — more so in a file that
        will be read by people who never saw the screen's own caveat.
        """
        _result, records, state, analysis = briefing_bundle
        report = insights_api._derivation_for(records[0], analysis)
        text = _docx_text(build_derivation_docx(report))
        assert "Numeric grounding" in text
        assert state.state_id in text

    def test_it_explains_the_method_rather_than_only_the_result(self, briefing_bundle):
        """
        The complaint this answers is that the derivation read as a black box:
        the steps and the figures were shown and nothing said where either
        came from, or whether a language model had a hand in them.
        """
        _result, records, _state, analysis = briefing_bundle
        report = insights_api._derivation_for(records[0], analysis)
        assert "deterministic" in report.method
        assert "checked back against the computed results" in report.method


class TestNothingOnScreenIsAStorageKey:
    @pytest.mark.parametrize("key,expected", [
        ("n_facilities_open", "Sites open"),
        ("n_facilities_closed", "Sites not used"),
        ("avg_utilization_pct", "Average utilisation"),
        ("max_utilization_pct", "Busiest site"),
        ("pct_demand_in_sla", "Demand within its lead time"),
        ("total_carbon_kg", "Transport emissions"),
        ("business_network_cost", "Total network cost"),
        ("demand_fill_rate", "Demand met"),
    ])
    def test_the_metrics_a_reader_sees_are_named_in_english(self, key, expected):
        assert metric_label(key) == expected

    def test_an_unmapped_key_still_reads_as_it_did(self):
        """A list of corrections, not a registry to keep in step."""
        assert metric_label("some_new_metric") == "Some New Metric"

    def test_a_fill_rate_is_a_percentage_not_a_ratio(self):
        """
        "The demand fill rate is 1.000" is the storage format. It was the
        headline figure of the service finding on the Overview tile and on
        every Insights row.
        """
        display, unit = _display(1.0, "demand_fill_rate")
        assert display == "100.0%"
        # The UNIT stays a ratio on purpose: `value` is still 1.0 and it is
        # what a chart plots, so calling this "percent" would let an axis
        # group it with figures that really are out of 100.
        assert unit == "ratio"

    def test_an_index_is_not_turned_into_a_percentage(self):
        """
        An REI of 1.02 is a score, not a proportion of anything, and "102%"
        would assert a whole it does not have.
        """
        display, unit = _display(1.02, "rei")
        assert display == "1.020"
        assert unit == "ratio"

    def test_a_count_of_sites_has_no_decimal_places(self):
        display, unit = _display(5, "n_facilities_open")
        assert display == "5"
        assert unit == "count"

    def test_money_in_the_prose_carries_its_currency(self, briefing_bundle):
        """
        The cost narrative read "business network cost at 150,627.70" beside
        an evidence chip reading ₹150,627.70 — the same figure twice on one
        card, once with its unit and once without. On a USD network it was a
        bare quantity in no unit at all.
        """
        _result, records, _state, _analysis = briefing_bundle
        cost = next(r for r in records if r["theme"] == "Cost")
        assert "₹" in cost["narrative"], cost["narrative"]

    def test_no_finding_quotes_a_bare_ratio_at_the_reader(self, briefing_bundle):
        """
        The sweep the specific cases above are examples of: nothing a reader
        sees may be a three-decimal proportion or a Title-Cased storage key.
        """
        _result, records, _state, _analysis = briefing_bundle
        for record in records:
            assert "1.000" not in record["narrative"], record["narrative"]
            for row in record["evidence"]:
                assert not row["label"].startswith("N "), row
                assert "Pct" not in row["label"], row
                assert "Kg" not in row["label"], row


class TestTheDocumentAndTheScreenCannotDisagree:
    """
    The bug this closes. `GET /api/insights` is CACHED per network version —
    a briefing is derived data about one version of one network, and
    recomputing it per dashboard load would pay for a reasoning pass and, on a
    fresh process, a solve. The document route was not cached: it resolved the
    twin state and ran its own reasoning pass.

    So the list a reader was looking at and the list the download searched
    were two different computations. They agreed until a hydration published a
    fresher state, at which point the reasoning pass wrote a slightly
    different headline — and the insight id is a digest of the headline. The
    reader opened a finding, pressed Download, and was told "'INS_NETWORK_…'
    is not a finding on this network's current analysis" about the finding
    filling their screen.
    """

    def test_a_record_carries_the_scope_it_was_computed_in(self, briefing_bundle):
        """
        Without it the client had to assume, and it assumed NETWORK — so a
        facility-scoped finding, which the deep dive opens exactly as readily,
        could never be downloaded at all.
        """
        _result, records, _state, _analysis = briefing_bundle
        assert all(r["scope"] == "NETWORK" for r in records), records[0]
        assert all("entity_id" in r for r in records)

    def test_the_route_reads_the_cached_briefing_rather_than_its_own(self):
        import inspect

        source = inspect.getsource(insights_api.create_insights_blueprint)
        body = source[source.index("def insight_document"):]
        body = body[:body.index("@bp.errorhandler")]
        assert "_briefing_analysis(" in body, body
        # And no second reasoning pass, which is what made them diverge.
        assert "_briefing_for(" not in body, body
        assert "_resolve_state(" not in body, body

    def test_it_accepts_the_scope_the_record_states(self):
        import inspect

        source = inspect.getsource(insights_api.create_insights_blueprint)
        body = source[source.index("def insight_document"):]
        body = body[:body.index("@bp.errorhandler")]
        assert 'request.args.get("scope")' in body
        assert 'request.args.get("entity_id")' in body

    def test_the_client_sends_the_records_own_scope(self):
        from pathlib import Path

        root = Path(insights_api.__file__).resolve().parents[3] / "app" / "frontend"
        service = (root / "js" / "integration" / "services"
                   / "insight-service.js").read_text(encoding="utf-8")
        block = service[service.index("async downloadDerivation("):]
        block = block[:block.index("\n  },")]
        assert "options.scope" in block, block
        assert "entity_id" in block, block

        detail = (root / "js" / "insight-detail.js").read_text(encoding="utf-8")
        call = detail[detail.index("insightService.downloadDerivation("):]
        call = call[:call.index("});")]
        assert "record.scope" in call, call
