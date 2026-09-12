"""
The derivation report, and its .docx rendering.

STRUCTURE, NOT PROSE
--------------------
`DerivationReport` is a description of one conclusion and the route to it:

    subject        what the report is about ("Capacity · whole network")
    conclusion     the finding, in the engine's own words
    summary        the prose explaining it, also the engine's
    method          how a reader should understand the derivation
    steps          the working, in the order it was done
    assumptions    what the model was told to take as given
    limitations    what it could not establish
    provenance     which run this came from, and whether it was checked

Every string arrives already written and already formatted. Nothing here
computes, rounds, re-words or infers: a document that recomputed a figure
would be a second, unverified engine whose output carries a letterhead and
gets forwarded to people who will never see this screen.

WHY .docx AND NOT PDF
---------------------
Because the document is a starting point. A planner takes it into a deck, cuts
two sections, adds their own context and sends it on — and a PDF makes them
retype it. `python-docx` writes a file Word and Google Docs both open and
neither complains about.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import List, Optional, Sequence

# Kearney purple, as the rest of the product uses it.
_ACCENT = (0x6B, 0x2F, 0xA0)
_MUTED = (0x6B, 0x72, 0x80)


@dataclass(frozen=True)
class Figure:
    """One number the conclusion rests on, already formatted.

    `value` is the display string — "97.20%", "₹1,435,985.00", "Not available"
    — never a float. The engine that computed it decided how it reads, and a
    second opinion about that in this file is how a document and the screen it
    came from end up disagreeing about the same number.
    """

    label: str
    value: str
    #: What part this figure played: the measurement, what it was compared
    #: against, or a driver behind it.
    role: str = "Measured"
    #: Which engine computed it.
    source: str = ""


@dataclass(frozen=True)
class Variable:
    """
    One symbol in an equation: what it means, what it was, where it came from.

    A formula with undefined symbols is decoration. `value` is the display
    string for THIS network — the same string the screen shows — so a reader
    can follow the arithmetic with the actual numbers rather than being shown
    the shape of a calculation and left to trust it.
    """

    symbol: str
    meaning: str
    value: str = ""
    source: str = ""


@dataclass(frozen=True)
class Equation:
    """
    One calculation, three ways: symbolically, with its symbols defined, and
    with this network's numbers substituted in.

    WHY ALL THREE. The symbolic form is what a reader checks against their own
    understanding of the metric. The substitution is what they check against
    the figure on the screen. Neither alone is convincing: a formula with no
    numbers cannot be verified, and an arithmetic line with no formula cannot
    be generalised to the next network.

    `substituted` is written by the caller from the same values the figures
    table carries, never re-derived here — a document whose worked example
    disagrees with its own results table is worse than one with no example.
    """

    name: str
    formula: str
    variables: Sequence[Variable] = field(default_factory=tuple)
    substituted: str = ""
    note: str = ""


@dataclass(frozen=True)
class DerivationStep:
    """One stage of the working, with the figures read at that stage."""

    title: str
    detail: str = ""
    figures: Sequence[Figure] = field(default_factory=tuple)
    #: The arithmetic behind this stage. Empty where the stage is a reading
    #: rather than a calculation — "the solver reported these flows" has no
    #: equation, and inventing one for it would be describing work that was
    #: not done.
    equations: Sequence[Equation] = field(default_factory=tuple)


@dataclass
class DerivationReport:
    subject: str
    conclusion: str
    summary: str = ""
    method: str = ""
    steps: List[DerivationStep] = field(default_factory=list)
    recommended_action: str = ""
    assumptions: List[str] = field(default_factory=list)
    limitations: List[str] = field(default_factory=list)
    #: "Source: … (state_id)". One line, written by the caller.
    provenance: str = ""
    #: An explanation of the working in plain English, and one line saying who
    #: wrote it. Built by `netgravity.reporting.narration`, which verifies
    #: every figure in it against the figures below before it gets here — see
    #: that module for why a model is allowed to write this at all.
    #:
    #: Optional in the strongest sense: a report with none is a complete
    #: report, and that is what one looks like whenever the gateway is
    #: unconfigured, over budget, or returned something that failed the check.
    narrative: List[str] = field(default_factory=list)
    #: Who wrote `narrative`, in a sentence a reader can act on. Printed even
    #: when the narrative itself is empty, because "this was withheld and why"
    #: is a fact the reader of a derivation is entitled to.
    narrative_note: str = ""
    #: Free-form label for the kind of thing this describes — "Insight",
    #: "Demand forecast". Printed above the title.
    kind: str = "Analysis"
    #: When this was produced, already formatted for a reader.
    generated_at: str = ""

    def filename(self) -> str:
        """A filename a person can find again in a downloads folder."""
        import re

        stem = re.sub(r"[^A-Za-z0-9]+", "-", f"{self.kind} {self.subject}").strip("-")
        return f"NetGravity-{stem or 'derivation'}.docx"


def _shade(cell, hex_colour: str) -> None:
    """Fill one table cell. python-docx has no API for it; this is the XML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour)
    tc_pr.append(shd)


def _para(doc, text: str, *, size: int = 10, bold: bool = False,
          colour: Optional[tuple] = None, space_after: int = 6,
          italic: bool = False):
    from docx.shared import Pt, RGBColor

    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if colour is not None:
        run.font.color.rgb = RGBColor(*colour)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _heading(doc, text: str) -> None:
    from docx.shared import Pt

    p = _para(doc, text.upper(), size=9, bold=True, colour=_ACCENT, space_after=4)
    p.paragraph_format.space_before = Pt(14)


def _equation_block(doc, equation: "Equation") -> None:
    """
    One equation: its name, the formula, what each symbol is, and the same
    formula with this network's numbers in it.

    The formula is set in a monospaced face and indented, because a reader
    scanning for "what was actually calculated" finds it by shape before they
    read a word of it.
    """
    from docx.shared import Pt, Inches

    if equation.name:
        _para(doc, equation.name, size=9.5, bold=True, space_after=2)

    formula = _para(doc, equation.formula, size=10.5, space_after=2)
    formula.paragraph_format.left_indent = Inches(0.25)
    for run in formula.runs:
        run.font.name = "Consolas"

    if equation.substituted:
        worked = _para(doc, equation.substituted, size=10.5, space_after=4,
                       colour=_ACCENT)
        worked.paragraph_format.left_indent = Inches(0.25)
        for run in worked.runs:
            run.font.name = "Consolas"
            run.bold = True

    variables = [v for v in (equation.variables or ()) if v is not None]
    if variables:
        table = doc.add_table(rows=1 + len(variables), cols=4)
        table.style = "Table Grid"
        table.autofit = False
        widths = (Inches(0.85), Inches(2.7), Inches(1.5), Inches(1.75))
        for col, (head, width) in enumerate(
                zip(("Symbol", "What it is", "This network", "From"), widths)):
            cell = table.cell(0, col)
            cell.text = ""
            run = cell.paragraphs[0].add_run(head)
            run.bold = True
            run.font.size = Pt(8.5)
            _shade(cell, "F3ECFA")
            cell.width = width
        for r, variable in enumerate(variables, start=1):
            for col, text in enumerate((variable.symbol, variable.meaning,
                                        variable.value, variable.source)):
                cell = table.cell(r, col)
                cell.text = ""
                run = cell.paragraphs[0].add_run(text or "—")
                run.font.size = Pt(8.5)
                if col == 0:
                    run.font.name = "Consolas"
                    run.bold = True
                cell.width = widths[col]
        doc.add_paragraph()

    if equation.note:
        _para(doc, equation.note, size=9, colour=_MUTED, italic=True,
              space_after=6)


def build_derivation_docx(report: DerivationReport) -> bytes:
    """
    Render one report as a .docx, and return its bytes.

    Sections that carry nothing are omitted rather than printed empty: a
    document with an "Assumptions" heading over a blank suggests the model
    made none, which is a claim, and a different one from "none were
    recorded".
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches

    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # ── Title block ──────────────────────────────────────────────────
    _para(doc, report.kind.upper(), size=8, bold=True, colour=_ACCENT,
          space_after=2)
    title = _para(doc, report.conclusion, size=16, bold=True, space_after=4)
    title.paragraph_format.space_before = Pt(0)
    _para(doc, report.subject, size=10, colour=_MUTED, space_after=2)
    if report.generated_at:
        _para(doc, report.generated_at, size=8.5, colour=_MUTED, space_after=10)

    if report.summary:
        _para(doc, report.summary, size=10.5, space_after=8)

    # ── How to read the derivation ───────────────────────────────────
    if report.method:
        _heading(doc, "How this was reached")
        _para(doc, report.method, size=10)

    # ── The same thing, joined up ────────────────────────────────────
    #
    # ABOVE the working rather than after it, because this is the part a
    # reader reads: the tables are what they come back to when they want to
    # check it. The attribution line sits immediately under the passage, not
    # in a footnote, so nobody can quote the paragraph without the sentence
    # saying where it came from.
    # A heading over nothing is a claim, and the wrong one: it suggests the
    # working could not be explained. When the gateway is simply absent the
    # section does not exist — the document is complete without it. It DOES
    # appear when a passage was written and withheld, because that is a fact
    # about this document the reader is entitled to.
    if report.narrative or report.narrative_note:
        _heading(doc, "The working, in plain terms")
        for paragraph in report.narrative:
            _para(doc, paragraph, size=10.5, space_after=8)
        if report.narrative_note:
            _para(doc, report.narrative_note, size=8.5, colour=_MUTED,
                  italic=True, space_after=4)

    # ── The working ──────────────────────────────────────────────────
    for index, step in enumerate(report.steps, start=1):
        _heading(doc, f"Step {index} — {step.title}")
        if step.detail:
            _para(doc, step.detail, size=10, space_after=6)
        # ── the arithmetic, before the figures it produced ──────────
        #
        # A reader following a derivation wants the rule before the result:
        # the equation says what was computed, the substitution shows it being
        # computed, and the table below records what came out.
        for equation in (step.equations or ()):
            _equation_block(doc, equation)

        rows = [f for f in step.figures if f is not None]
        if not rows:
            continue

        table = doc.add_table(rows=1 + len(rows), cols=4)
        table.style = "Table Grid"
        table.autofit = False
        widths = (Inches(2.6), Inches(1.5), Inches(1.3), Inches(1.4))
        headers = ("Figure", "Value", "Role", "Computed by")
        for col, (head, width) in enumerate(zip(headers, widths)):
            cell = table.cell(0, col)
            cell.text = ""
            run = cell.paragraphs[0].add_run(head)
            run.bold = True
            run.font.size = Pt(9)
            _shade(cell, "F3ECFA")
            cell.width = width
        for r, figure in enumerate(rows, start=1):
            for col, text in enumerate((figure.label, figure.value,
                                        figure.role, figure.source)):
                cell = table.cell(r, col)
                cell.text = ""
                run = cell.paragraphs[0].add_run(text or "—")
                run.font.size = Pt(9)
                if col == 1:
                    run.bold = True
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                cell.width = widths[col]
        doc.add_paragraph()

    # ── What follows from it ─────────────────────────────────────────
    if report.recommended_action:
        _heading(doc, "Recommended action")
        _para(doc, report.recommended_action, size=10.5)

    if report.assumptions:
        _heading(doc, "What the model was given")
        for line in report.assumptions:
            p = doc.add_paragraph(line, style="List Bullet")
            p.runs[0].font.size = Pt(9.5)

    if report.limitations:
        _heading(doc, "What this does not establish")
        for line in report.limitations:
            p = doc.add_paragraph(line, style="List Bullet")
            p.runs[0].font.size = Pt(9.5)

    if report.provenance:
        _heading(doc, "Provenance")
        _para(doc, report.provenance, size=8.5, colour=_MUTED)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
