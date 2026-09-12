"""Word and JSON representations of the same authenticated calculation report."""
from io import BytesIO
import json
import re


def text(value):
    if value is None:
        return "Not available"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return str(value)


def leaves(value, prefix=""):
    if isinstance(value, dict) and value:
        for key, child in value.items():
            yield from leaves(child, f"{prefix} / {key}".strip(" /"))
    elif isinstance(value, list) and value:
        for index, child in enumerate(value, 1):
            yield from leaves(child, f"{prefix} / record {index}".strip(" /"))
    else:
        yield prefix or "Value", "None recorded" if value in ({}, []) else text(value)


def build_word(report):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(26)
    doc.styles["Heading 1"].font.size = Pt(17)
    doc.styles["Heading 2"].font.size = Pt(13)
    footer = section.footer.paragraphs[0]
    footer.add_run("NetGravity calculation audit  ·  Page ").font.size = Pt(9)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)

    def table(rows, headers=("Calculation field", "Recorded value")):
        rows = list(rows)
        if not rows:
            doc.add_paragraph("No records were captured for this section.")
            return
        tab = doc.add_table(rows=1, cols=len(headers))
        tab.autofit = False
        for column in tab.columns:
            column.width = Inches(7 / len(headers))
        for cell, label in zip(tab.rows[0].cells, headers):
            cell.text = str(label).replace("_", " ")
        repeat = OxmlElement("w:tblHeader")
        tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            cells = tab.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = text(value)
        for i, row in enumerate(tab.rows):
            if all(len(cell.text) < 1000 for cell in row.cells):
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "233248" if i == 0 else ("F1F4F8" if i % 2 else "FFFFFF"))
                props.append(shade)
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + edge)
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        border.set(qn("w:" + key), value)
                    borders.append(border)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        doc.add_paragraph()

    def source_table(value, path=""):
        """Keep repeated data tabular and vectors compact, without dropping values."""
        primitive = lambda item: not isinstance(item, (dict, list))
        if isinstance(value, list) and value:
            if all(primitive(item) for item in value):
                table([(path or "Values in order", json.dumps(value, ensure_ascii=False))])
                return
            if all(isinstance(item, dict) and item for item in value):
                keys = list(dict.fromkeys(key for item in value for key in item))
                if len(keys) <= 7 and all(primitive(v) for item in value for v in item.values()):
                    table([[item.get(key) for key in keys] for item in value], headers=keys)
                    return
            for index, child in enumerate(value, 1):
                name = f"{path} / Record {index}".strip(" / ")
                paragraph = doc.add_paragraph(name)
                paragraph.runs[0].bold = True
                paragraph.paragraph_format.keep_with_next = True
                source_table(child)
        elif isinstance(value, dict) and value:
            simple = [(key.replace("_", " "), val) for key, val in value.items() if primitive(val)]
            if simple:
                table(simple)
            for key, child in value.items():
                if primitive(child):
                    continue
                name = f"{path} / {key.replace('_', ' ')}".strip(" / ")
                paragraph = doc.add_paragraph(name)
                paragraph.runs[0].bold = True
                paragraph.paragraph_format.keep_with_next = True
                source_table(child)
        else:
            table(leaves(value, path))

    title = re.sub(r"[^\w\s]", "", report.get("title", "Calculation report"))
    doc.add_paragraph(title, "Title")
    doc.add_paragraph("This report explains the selected analytical run, its calculation methods and the recorded numbers behind the results. Use the worked calculations to understand the outcome and the source records to reproduce or challenge its assumptions.")
    doc.add_heading("1 Result and scope", level=1)
    table(leaves({key: report.get(key) for key in ("project_id", "snapshot_id", "execution_id", "computed_at", "scope")}))
    doc.add_heading("2 Methodology and assumptions", level=1)
    for paragraph in report.get("methodology", []):
        doc.add_paragraph(paragraph)
    for limitation in report.get("limitations", []):
        doc.add_paragraph(limitation)
    doc.add_heading("3 Worked calculations", level=1)
    for row in report.get("metrics", []):
        doc.add_heading(re.sub(r"[^\w\s]", "", row["label"]), level=2)
        doc.add_paragraph(f"Result: {text(row['value'])} {row.get('unit') or ''}. Status: {row['status']}.")
        doc.add_paragraph(row["formula"])
        if row.get("worked"):
            doc.add_paragraph(row["worked"])
        if row.get("note"):
            doc.add_paragraph(row["note"])
        source_table(row.get("inputs"))
    doc.add_heading("4 Source data and calculation records", level=1)
    doc.add_paragraph("The following records preserve the input values and actual engine output for this run. Numbered records identify array positions; entity IDs and period fields provide business traceability. Values are not rounded again for export.")
    for key, source in report.get("sources", {}).items():
        doc.add_heading(re.sub(r"[^\w\s]", "", key.replace("_", " ").capitalize()), level=2)
        source_table(source)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def calculation_response(report):
    from flask import jsonify, request, send_file
    if request.args.get("format") == "docx":
        response = send_file(build_word(report), as_attachment=True,
                             download_name=f"netgravity-{report['kind']}-calculations.docx",
                             mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    else:
        response = jsonify(report)
    response.headers["Cache-Control"] = "private, no-store"
    return response
