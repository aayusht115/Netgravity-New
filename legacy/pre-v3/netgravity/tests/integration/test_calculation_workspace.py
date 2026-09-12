"""Real-engine calculation trails, selected-scenario scope and AI fail-closed behavior."""
import json
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from flask import Flask

from app.backend.services.calculations import network_report, forecast_report
from app.backend.services.executive_briefing import scenario_briefing
from netgravity.tests.integration.test_warehouse_deep_dive import _network, _solve
from netgravity.orchestrator.metrics.registry import KPIRegistry
from netgravity.orchestrator.core.execution_context import ExecutionContext
from netgravity.forecasting import DemandPoint, DemandTimeSeries, ForecastingService, ForecastRequest


def test_cost_inputs_recompose_authoritative_total_and_exclude_penalty():
    raw, state = _solve(_network())
    ctx = ExecutionContext(execution_id="audit-test", network_states={"optimization.solve": state})
    kpis = {k:v.model_dump(mode="json") for k,v in KPIRegistry().network_kpis(ctx).items()}
    report = network_report({"calculation_state": state.model_dump(mode="json"), "calculation_trace": state.calculation_trace,
                             "kpis": kpis, "execution_id": ctx.execution_id}, project_id="owned", snapshot_id="snap")
    total = report["metrics"][0]
    assert sum(c["value"] for c in total["inputs"] if c["included"]) == pytest.approx(total["value"], abs=1e-4)
    assert total["reconciliation_difference"] == pytest.approx(0, abs=1e-4)
    assert not next(c for c in total["inputs"] if c["id"] == "shortage_cost")["included"]
    assert state.calculation_trace["solver_output"]["flow_decisions"] == [row.model_dump(mode="json") for row in raw.flow_decisions]
    assert "calculation_trace" not in state.model_dump(mode="json"), "Raw audit input must not inflate ordinary reasoning payloads"
    assert state.calculation_trace["engine_parameters"]["modelled_periods"]
    assert "calculation_parameters" not in raw.model_dump(mode="json")
    transport = next(m for m in report["metrics"] if m["id"] == "transport_cost")
    assert sum(f["flow_units"] * f["rate_per_unit"] for f in transport["inputs"]) == pytest.approx(transport["value"], abs=.001)


def test_quantile_steps_and_backtests_reproduce_each_reported_number():
    from app.backend.api.forecast import _serialise_series
    history = [DemandPoint(period=i+1,quantity=100+i*7+(i%4)*3) for i in range(24)]
    sf = ForecastingService().forecast(ForecastRequest(snapshot_id="snap",horizon=6,run_backtest=True,
        series=[DemandTimeSeries(market_id="M",product_id="P",history=history)])).series[0]
    diag = sf.calculation_trace["diagnostics"]
    for step, point in zip(diag["forecast_steps"], sf.points):
        values = [max(0,sum(x*b for x,b in zip(step["features"],diag["coefficients"][q]))) for q in ("p10","p50","p90")]
        assert sorted(values) == pytest.approx([point.p10,point.p50,point.p90])
    audit = sf.accuracy.calculation_trace
    assert audit["total_absolute_error"] / audit["total_actual"] == pytest.approx(sf.accuracy.wape,abs=1e-6)
    assert all(row["absolute_error"] == abs(row["actual"]-row["prediction"]) for row in audit["folds"])
    row = {**_serialise_series(sf), "history":[p.model_dump(mode="json") for p in history]}
    report = forecast_report(row,project_id="owned",snapshot_id="snap",execution_id="run",horizon=6)
    assert report["metrics"][0]["value"] == sum(p.mean for p in sf.points)
    assert report["metrics"][1]["value"] == pytest.approx((sum(p.mean for p in sf.points)/sum(p.quantity for p in history[-6:])-1)*100)


def test_inventory_coefficients_are_the_ones_used_by_the_actual_solve():
    raw,state=_solve(_network(inventory=True))
    parameters=state.calculation_trace["engine_parameters"]
    assert parameters["inventory_coefficients"]
    by_pair={(c["facility_id"],c["market_id"]):c for c in parameters["inventory_coefficients"]}
    reconstructed=sum(by_pair.get((f.origin_id,f.destination_id),{}).get("unit_inv_cost_by_product",{}).get(f.product_id,0)*f.flow_units
                      for f in raw.flow_decisions)
    assert reconstructed == pytest.approx(raw.objective_components["inventory_cost"],abs=.001)


def test_forecast_zero_history_does_not_invent_growth_or_backtests():
    report = forecast_report({"market_id":"M","product_id":"P","status":"OK",
        "points":[{"period":1,"mean":0,"p10":0,"p90":0}], "history":[{"period":1,"quantity":0}]},
        project_id="owned",snapshot_id="s",execution_id="e",horizon=1)
    assert report["metrics"][0]["value"] == 0
    assert report["metrics"][1]["value"] is None
    assert report["metrics"][4]["value"] is None


def _record():
    def k(v): return {"value":v,"status":"VALID","unit":"units"}
    return {"id":"SCN_A","name":"North demand growth","project_id":"owned","snapshot_id":"snap","execution_id":"run",
            "request":{"action":"CHANGE_DEMAND","demand_multiplier":1.1},"feasible":True,
            "baseline_kpis":{"business_network_cost":k(100),"unserved_demand":k(0),"max_utilization_pct":k(75)},
            "scenario_kpis":{"business_network_cost":k(115),"unserved_demand":k(5),"max_utilization_pct":k(99)}}


def test_ai_unavailable_does_not_pose_fixed_advice_as_generated():
    response = scenario_briefing(_record(),gateway=SimpleNamespace(available=False))
    assert response["status"] == "AI_UNAVAILABLE"
    assert response["recommendations"] == []
    assert len(response["findings"]) == 3
    assert response["findings"][0]["id"] == "unserved_demand"
    assert next(f for f in response["findings"] if f["id"] == "business_network_cost")["scenario"] == 115
    assert "10% increase in demand" in response["summary"]["description"]


@pytest.mark.parametrize("refs,text", [(["made-up"],"Add capacity"),(["unserved_demand"],"Save 400 percent")])
def test_unverified_model_advice_is_rejected(refs,text):
    gateway=SimpleNamespace(available=True,generate=lambda *a,**k:SimpleNamespace(output=json.dumps({
        "summary":"Capacity is constrained.","recommendations":[{"text":text,"evidence_ids":refs}],"finding_ids":[]}),model_name="gpt-5.6-luna"))
    response=scenario_briefing(_record(),gateway=gateway)
    assert response["status"] == "AI_FAILED"
    assert response["recommendations"] == []


def test_model_advice_carries_exact_selected_scenario_and_evidence():
    gateway=SimpleNamespace(available=True,generate=lambda *a,**k:SimpleNamespace(output=json.dumps({
        "summary":"Additional demand exceeds available capacity.","recommendations":[
            {"text":"Test capacity expansion before committing to the higher demand plan.","evidence_ids":["unserved_demand"]}],
        "finding_ids":["unserved_demand","business_network_cost","max_utilization_pct"]}),model_name="gpt-5.6-luna"))
    response=scenario_briefing(_record(),gateway=gateway)
    assert response["scenario_id"] == "SCN_A"
    assert response["status"] == "READY"
    assert response["model"] == "gpt-5.6-luna"
    assert response["findings"][0]["id"] == "unserved_demand"


def test_unverified_summary_numbers_are_rejected():
    gateway=SimpleNamespace(available=True,generate=lambda *a,**k:SimpleNamespace(output=json.dumps({
        "summary":"This scenario saves 500 dollars.","recommendations":[
            {"text":"Test capacity expansion.","evidence_ids":["unserved_demand"]}],
        "finding_ids":[]}),model_name="gpt-5.6-luna"))
    assert scenario_briefing(_record(), gateway=gateway)["status"] == "AI_FAILED"


def test_transient_advice_failure_is_not_persisted(monkeypatch):
    from app.backend.services.analysis_store import AnalysisService
    monkeypatch.setattr("app.backend.services.persistence.load_analysis", lambda *args: None)
    save = Mock()
    monkeypatch.setattr("app.backend.services.persistence.save_analysis", save)
    service=AnalysisService()
    compute=Mock(side_effect=[{"briefing":{"status":"AI_FAILED"}}, {"briefing":{"status":"READY","snapshot_id":"snap"}}])
    def fetch():
        return service.get("snap","data-version",compute,variant="briefing-test",cache_if=lambda value:value["briefing"]["status"]=="READY")
    assert fetch()["briefing"]["status"] == "AI_FAILED"
    assert not save.called
    assert fetch()["briefing"]["snapshot_id"] == "snap"
    assert fetch()["briefing"]["status"] == "READY"
    assert compute.call_count == 2
    assert save.call_count == 1


def test_word_export_preserves_scope_sections_and_full_precision():
    from io import BytesIO
    from docx import Document
    from app.backend.services.calculation_export import build_word
    report=forecast_report({"market_id":"M","product_id":"P","status":"OK","history":[{"quantity":100}],
        "points":[{"period":1,"mean":102.123456789,"p10":90,"p90":115}]},
        project_id="owned",snapshot_id="snap",execution_id="exact-run",horizon=1)
    doc=Document(BytesIO(build_word(report).getvalue()))
    content="\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "102.123456789" in content
    assert "exact-run" in content
    for title in ("1 Result and scope", "2 Methodology and assumptions", "3 Worked calculations", "4 Source data and calculation records"):
        assert title in content


def test_openai_transport_sends_requested_model_and_never_substitutes(monkeypatch):
    from netgravity.orchestrator.agents.openai_gateway import OpenAIResponsesGateway
    monkeypatch.setenv("OPENAI_API_KEY","not-a-real-key")
    monkeypatch.setenv("OPENAI_MODEL","gpt-5.6-luna")
    monkeypatch.delenv("NETGRAVITY_DISABLE_LLM",raising=False)
    post=Mock(return_value=SimpleNamespace(status_code=200,json=lambda:{"status":"completed","id":"response-test",
        "model":"gpt-5.6-luna","output":[{"type":"message","content":[{"type":"output_text","text":"Verified response"}]}]}))
    monkeypatch.setattr("requests.post",post)
    response=OpenAIResponsesGateway().generate("Test only")
    assert response.model_name == "gpt-5.6-luna"
    assert post.call_args.kwargs["json"]["model"] == "gpt-5.6-luna"
    assert post.call_args.kwargs["json"]["store"] is False


@pytest.fixture
def calculation_client(monkeypatch):
    from app.backend.api import forecast, scenarios, kpis
    from app.backend.services import security, analysis_store, persistence
    from app.backend.services.errors import NotFoundError, UnauthenticatedError
    from netgravity.orchestrator import build_orchestrator
    history=[DemandTimeSeries(market_id="MKT_N",product_id="P1",history=[DemandPoint(period=i+1,quantity=80+i*2) for i in range(12)])]
    orch=build_orchestrator(network=_network(periods=[200,220]),enable_llm=False,history_provider=lambda snapshot:(history,[]))
    def session(token):
        if not token: raise UnauthenticatedError("Sign in required")
        return SimpleNamespace(user_id=token)
    monkeypatch.setattr(security.auth_service,"resolve_session",session)
    def owned(project_id,*,user_id):
        if project_id != "owned" or user_id != "owner": raise NotFoundError("Project not found")
        return orch.snapshots.current_id
    monkeypatch.setattr(forecast.project_registry,"snapshot_for",owned)
    monkeypatch.setattr(forecast.demand_history_store,"for_snapshot",lambda snapshot:(history,[]))
    monkeypatch.setattr(forecast,"_uploaded_signals_for",lambda *a:([],[]))
    monkeypatch.setattr(analysis_store,"analysis_service",analysis_store.AnalysisService())
    monkeypatch.setattr(kpis,"analysis_service",analysis_store.analysis_service)
    monkeypatch.setattr(persistence,"load_scenarios",lambda:[])
    monkeypatch.setattr(scenarios,"explanations_llm_enabled",lambda:False)
    app=Flask(__name__)
    app.register_blueprint(forecast.create_forecast_blueprint(orch))
    app.register_blueprint(scenarios.create_scenario_blueprint(orch))
    app.register_blueprint(kpis.create_kpi_blueprint(orch))
    app.config["TESTING"]=True
    with app.test_client() as client: yield client,orch


def test_forecast_calculations_are_bound_to_saved_run_and_authenticated_owner(calculation_client):
    client,orch=calculation_client
    headers={"Authorization":"Bearer owner"}
    payload=client.get("/api/forecast?project_id=owned&horizon=3",headers=headers).get_json()
    path=payload["calculation_path"]+"?project_id=owned&series_id=MKT_N/P1"
    report=client.get(path,headers=headers)
    assert report.status_code == 200,report.get_json()
    body=report.get_json()
    assert body["execution_id"] == payload["execution_id"]
    assert body["metrics"][0]["value"] == sum(p["mean"] for p in payload["series"][0]["points"])
    assert client.get(path).status_code == 401
    assert client.get(path,headers={"Authorization":"Bearer other"}).status_code == 404
    word=client.get(path+"&format=docx",headers=headers)
    assert word.status_code == 200
    assert word.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert word.data.startswith(b"PK")
    assert word.headers["Cache-Control"] == "private, no-store"


def test_scenario_audit_preserves_both_solved_sides(calculation_client):
    client,orch=calculation_client
    headers={"Authorization":"Bearer owner"}
    response=client.post("/api/scenarios/simulate",headers=headers,json={"project_id":"owned","name":"Demand growth test","action":"CHANGE_DEMAND","demand_multiplier":1.1})
    record=response.get_json()
    assert response.status_code == 201,record
    path=f"/api/scenarios/{record['id']}/calculations?project_id=owned"
    response=client.get(path,headers=headers)
    assert response.status_code == 200,response.get_json()
    report=response.get_json()
    assert report["execution_id"] == record["execution_id"]
    assert report["sources"]["scenario"]["input_and_solver_trace"]["solver_output"]["flow_decisions"]
    assert record["cost_components"]["scenario"]
    assert client.get(path,headers={"Authorization":"Bearer other"}).status_code == 404
    briefing=client.get(f"/api/scenarios/{record['id']}/briefing?project_id=owned",headers=headers).get_json()
    assert briefing["scenario_id"] == record["id"]
