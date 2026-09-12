"""
The scenario, as the document a decision gets taken from — and the
recommendations that document and the screen both read.

Three things are held here.

  * **What to do is decided once, on the server.** The list used to be derived
    in a browser render function from the same capacity block the backend
    already had, so the reasoning behind a recommendation could not be audited
    and had to be written a second time for the document. Two definitions of
    what this application recommends, free to disagree about a client's
    network.

  * **A recommendation is a network intervention.** Add capacity, reopen a
    site, build one, scope the growth, get the missing input. "Review the
    proposed changes" was the first item on every scenario ever solved: it
    tells a reader to look at the screen they are already looking at, and it
    pushed the answers below itself.

  * **Nothing in the file is gibberish.** A cost with no currency, a fill rate
    printed as the ratio 1.000, and a primary key inside a sentence — "Proceed
    with a detailed implementation review of Key probe SCN_084df97a" — all
    reached a document written for senior leadership.
"""

from __future__ import annotations

import io
import re
import uuid
from pathlib import Path

import pytest

import app.backend.api.scenarios as scenarios_api


_FRONTEND = Path(scenarios_api.__file__).resolve().parents[3] / "app" / "frontend"


def _asset(*parts) -> str:
    return _FRONTEND.joinpath(*parts).read_text(encoding="utf-8")


def _kpi(value, unit=""):
    """A KPI row in the shape the registry stores one."""
    return {"value": value, "unit": unit, "status": "VALID"}


def _record(**over):
    """A solved scenario record, of the shape `/simulate` persists."""
    base = {
        "id": "SCN_test0001",
        "name": "Peak season +40%",
        "snapshot_id": "snap_abc123456789",
        "execution_id": "exe-1",
        "feasible": True,
        "request": {"action": "CHANGE_DEMAND", "demand_multiplier": 1.4},
        "overrides": ["CHANGE_DEMAND all x1.4 (8 rows)"],
        "scenario_kpis": {
            "business_network_cost": _kpi(173079.20, "INR"),
            "transport_cost": _kpi(59677.0, "INR"),
            "facility_cost": _kpi(95000.0, "INR"),
            "demand_fill_rate": _kpi(1.0, "fraction"),
            "unserved_demand": _kpi(0.0, "units"),
            "avg_utilization_pct": _kpi(73.08, "pct"),
            "n_facilities_open": _kpi(5, "count"),
        },
        "baseline_kpis": {
            "business_network_cost": _kpi(150627.70, "INR"),
            "transport_cost": _kpi(45890.0, "INR"),
            "facility_cost": _kpi(95000.0, "INR"),
            "demand_fill_rate": _kpi(1.0, "fraction"),
            "unserved_demand": _kpi(0.0, "units"),
            "avg_utilization_pct": _kpi(56.23, "pct"),
            "n_facilities_open": _kpi(5, "count"),
        },
        "capacity_response": {},
        "explanation": {"card": {"headline": "This change increases what the "
                                             "network costs",
                                 "details": ["Cost proximity to today"]}},
        "provenance": {"engine": "netgravity MILP (PuLP/HiGHS)",
                       "authoritative_source": "KPIRegistry"},
    }
    base.update(over)
    return base


def _text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# What to do about it
# ---------------------------------------------------------------------------
class TestARecommendationIsANetworkIntervention:
    def test_a_full_site_is_where_capacity_is_recommended(self):
        record = _record(capacity_response={
            "at_ceiling": [{"id": "DC_WEST", "name": "Western DC",
                            "util_pct": 100.0, "added_units": 800.0,
                            "region": "West"}]})
        actions = scenarios_api._recommended_actions(record)
        first = actions[0]
        assert first["key"] == "ADD_CAPACITY"
        assert "Western DC" in first["label"]
        # The REASON, which is what a leader is deciding on.
        assert "100% of its capacity" in first["reason"]
        assert "constraint" in first["reason"]

    def test_reopening_is_offered_before_building(self):
        """
        Capacity that already exists is cheaper to use than capacity that has
        to be built, so a plan that left a site closed while stranding demand
        is told to reopen it first.
        """
        record = _record(
            capacity_response={
                "idle": [{"id": "DC_NORTH", "name": "Northern DC",
                          "capacity": 4500.0, "region": "North"}],
                "at_ceiling": [{"id": "DC_WEST", "name": "Western DC",
                                "util_pct": 100.0, "region": "West"}]},
            scenario_kpis={**_record()["scenario_kpis"],
                           "unserved_demand": _kpi(1200.0, "units")})
        keys = [a["key"] for a in scenarios_api._recommended_actions(record)]
        assert keys.index("REOPEN_FACILITY") < keys.index("ADD_CAPACITY")

    def test_a_new_site_needs_a_region_with_nothing_left(self):
        """
        Only a region whose every site is full, with nothing closed to reopen
        and no headroom on anything open, honestly needs one built. Anything
        weaker recommends building where a reopening would have done.
        """
        without = scenarios_api._recommended_actions(_record(
            capacity_response={"at_ceiling": [{"id": "DC_WEST",
                                               "name": "Western DC",
                                               "util_pct": 100.0}]}))
        assert not any(a["key"] == "OPEN_NEW_FACILITY" for a in without)

        with_region = scenarios_api._recommended_actions(_record(
            capacity_response={
                "at_ceiling": [{"id": "DC_WEST", "name": "Western DC",
                                "util_pct": 100.0}],
                "regions_without_room": [{"region": "West"}]}))
        build = next(a for a in with_region if a["key"] == "OPEN_NEW_FACILITY")
        assert "West" in build["label"]
        assert "nowhere to go" in build["reason"]

    def test_high_risk_with_nothing_yet_full_still_gets_an_answer(self):
        """
        The gap this closes: the card reported "capacity risk: High" beside a
        site at 92.6% and recommended nothing about capacity, because the
        ceiling test had not tripped. A reader told the network is at risk and
        given no way to act on it has the worst of both.
        """
        record = _record(
            capacity_risk="HIGH",
            capacity_response={"at_ceiling": [], "working_harder": [
                {"id": "DC_WEST", "name": "Western DC", "util_pct": 92.6,
                 "region": "West"}]})
        actions = scenarios_api._recommended_actions(record)
        add = next(a for a in actions if a["key"] == "ADD_CAPACITY")
        assert "Western DC" in add["label"]
        assert "93% of its capacity" in add["reason"]
        assert "first site that will" in add["reason"]

    def test_nothing_to_recommend_is_stated_rather_than_left_blank(self):
        """
        A plan that fills nothing and strands nothing has no intervention to
        recommend. An empty "Recommended actions" heading answers nothing;
        this says what was found instead.
        """
        actions = scenarios_api._recommended_actions(_record(
            capacity_response={"at_ceiling": [], "idle": []},
            request={"action": "CHANGE_CAPACITY"}))
        assert len(actions) == 1
        assert actions[0]["key"] == "NO_ACTION"
        assert "serves all of the demand" in actions[0]["reason"]
        assert actions[0]["target"] == {}

    def test_every_action_carries_its_reason_and_its_rank(self):
        actions = scenarios_api._recommended_actions(_record(
            capacity_response={"at_ceiling": [{"id": "D", "name": "D",
                                               "util_pct": 99.0}]}))
        assert [a["priority"] for a in actions] == list(range(1, len(actions) + 1))
        for action in actions:
            assert action["reason"].strip()
            assert action["key"] in scenarios_api._ACTION_KEYS


# ---------------------------------------------------------------------------
# The document
# ---------------------------------------------------------------------------
class TestTheScenarioBecomesADocument:
    def test_it_answers_the_four_questions_a_decision_asks(self):
        report = scenarios_api._scenario_derivation(
            _record(), "Case-16", scenarios_api._recommended_actions(_record()))
        titles = [step.title for step in report.steps]
        assert "What was changed" in titles
        assert any("costs" in t for t in titles)
        assert any("service" in t for t in titles)
        assert any("recommended" in t for t in titles)

    def test_money_carries_the_currency_the_row_states(self):
        """
        "Peak season +40% demand costs 167,050.33 per period" is a document
        forwarded to people who cannot know from context whether that is
        rupees or dollars. The stored KPI row carries `unit: "INR"`, and it is
        the only place this network's currency is recorded on the record.
        """
        from netgravity.reporting import build_derivation_docx

        report = scenarios_api._scenario_derivation(_record(), "", [])
        text = _text(build_derivation_docx(report))
        assert "₹173,079" in text
        assert "₹150,628" in text
        assert "173,079 per period," not in text.replace("₹173,079", "")

    def test_a_fill_rate_is_a_percentage_not_a_ratio(self):
        from netgravity.reporting import build_derivation_docx

        report = scenarios_api._scenario_derivation(_record(), "", [])
        text = _text(build_derivation_docx(report))
        assert "100.0%" in text
        assert "1.000" not in text

    def test_the_largest_cost_line_is_in_the_table(self):
        """
        The table asked for `fixed_cost`; the KPI registry calls it
        `facility_cost`. So the largest single line in this network's cost —
        more than half the total — was missing while the components below it
        were listed, and the numbers shown did not add up to the total beside
        them.
        """
        from netgravity.reporting import build_derivation_docx

        report = scenarios_api._scenario_derivation(_record(), "", [])
        text = _text(build_derivation_docx(report))
        assert "Fixed facility" in text
        assert "₹95,000" in text

    def test_a_component_that_is_zero_on_both_sides_reads_as_unchanged(self):
        assert scenarios_api._delta_display(
            {"opening_cost": _kpi(0.0, "INR")},
            {"opening_cost": _kpi(0.0, "INR")}, "opening_cost") == "unchanged"
        assert scenarios_api._delta_display(
            {"c": _kpi(10.0, "INR")}, {"c": _kpi(0.0, "INR")},
            "c") == "nothing in the baseline to compare against"

    def test_the_recommendations_travel_with_it(self):
        from netgravity.reporting import build_derivation_docx

        record = _record(capacity_response={
            "at_ceiling": [{"id": "DC_WEST", "name": "Western DC",
                            "util_pct": 100.0, "region": "West"}]})
        actions = scenarios_api._recommended_actions(record)
        report = scenarios_api._scenario_derivation(record, "", actions)
        text = _text(build_derivation_docx(report))
        for action in actions:
            assert action["label"] in text
            assert action["reason"][:60] in text

    def test_it_says_a_scenario_is_not_a_decision(self):
        """
        Opening or closing a site is classified as a human decision by
        governance whatever the economics say. A document that reads as an
        approval is the one thing this must not be.
        """
        from netgravity.reporting import build_derivation_docx

        report = scenarios_api._scenario_derivation(_record(), "", [])
        text = _text(build_derivation_docx(report))
        assert "not a decision" in text
        assert "nothing in this document approves anything" in text

    def test_an_infeasible_scenario_says_so_in_its_conclusion(self):
        report = scenarios_api._scenario_derivation(
            _record(feasible=False), "", [])
        assert "no feasible plan" in report.conclusion


# ---------------------------------------------------------------------------
# Nothing on the card is a storage key
# ---------------------------------------------------------------------------
class TestNoInternalIdentifierReachesAReader:
    @pytest.mark.parametrize("before,after", [
        ("Proceed with a detailed implementation review of Key probe SCN_084df97a.",
         "Proceed with a detailed implementation review of Key probe."),
        ("The state tws_snap_9dd1b976fe01_optimized was used.",
         "The state was used."),
        ("Compare Peak season (SCN_432b7822) against the baseline.",
         "Compare Peak season against the baseline."),
    ])
    def test_a_record_id_is_removed_from_prose(self, before, after):
        from netgravity.orchestrator.reasoning.card import strip_internal_ids

        assert strip_internal_ids(before) == after

    def test_a_facility_id_is_not_an_internal_identifier(self):
        """
        `DC_CENTRAL` came out of the reader's own upload and is how they refer
        to the site themselves. Stripping it would remove the subject of the
        sentence.
        """
        from netgravity.orchestrator.reasoning.card import strip_internal_ids

        assert strip_internal_ids("DC_CENTRAL runs at 92%.") == "DC_CENTRAL runs at 92%."
        assert strip_internal_ids("PLANT_SOUTH carries the exposure.") \
            == "PLANT_SOUTH carries the exposure."

    def test_every_reader_facing_string_goes_through_it(self):
        from netgravity.orchestrator.reasoning.card import clean

        assert "SCN_084df97a" not in clean("Review Key probe SCN_084df97a now.")

    def test_the_saved_wording_is_invalidated_when_the_wording_changes(self):
        """
        A saved explanation is keyed by the RESULT it describes, which is right
        for the figures and says nothing about the sentences — so a correction
        to materially wrong prose changed no fingerprint and every project that
        had already been explained kept the older words for the life of the
        store.
        """
        from netgravity.orchestrator import explanations

        assert explanations.PROSE_VERSION >= 2
        first = explanations.fingerprint("exec-1", "v9")
        explanations.PROSE_VERSION += 1
        try:
            assert explanations.fingerprint("exec-1", "v9") != first
        finally:
            explanations.PROSE_VERSION -= 1


# ---------------------------------------------------------------------------
# The screen reads the server's list
# ---------------------------------------------------------------------------
class TestTheScreenRendersWhatTheServerDecided:
    def test_the_download_is_offered_on_the_card_and_in_the_drawer(self):
        js = _asset("js", "scenarios.js")
        assert "scn-download-doc" in js
        assert "scn-drawer-download" in js
        assert "async function downloadScenarioDerivation(button, scenarioId)" in js
        service = _asset("js", "integration", "services", "scenario-service.js")
        assert "/document" in service
        assert "CONFIG.DOCUMENT_TIMEOUT_MS" in service

    def test_the_drawer_ends_on_what_is_recommended(self):
        """
        The drawer is opened from "View full detail", so it is read by
        somebody who has seen the recommendation and wants the basis for it.
        Ending on a cost table left them to carry the recommendation across
        from the card in their head.
        """
        js = _asset("js", "scenarios.js")
        drawer = js[js.index("export function openScenarioDrawer("):]
        drawer = drawer[:drawer.index("\n// ─── Open Metric Drilldown")]
        assert "What is recommended, and why" in drawer
        assert "recommendedActions(scn, comparisonState.data)" in drawer

    def test_the_card_does_not_label_its_own_words_rule_based(self):
        """
        Jargon about this application's internals is not what a chip beside a
        recommendation is for — and it was showing on a build with a working
        gateway, because the model's reply was being truncated at the output
        budget and silently discarded. The label was reporting a defect as if
        it were a design.

        Nothing is claimed instead: calling template prose "AI" would be a
        false statement about provenance, and the honest account stays in
        "How this was decided".
        """
        js = _asset("js", "scenarios.js")
        head = js[js.index("function takeHeadHtml("):]
        head = head[:head.index("\n}\n")]
        # What the function RENDERS. The phrase survives in this file as a
        # comment recording why the chip changed, which is the opposite of the
        # defect and must not fail its own test.
        markup = re.sub(r"//[^\n]*", "", head)
        assert "Rule-based" not in markup, markup
        assert "scn-take-source" in markup

    def test_no_screen_names_the_template_to_a_reader(self):
        for asset in (("js", "scenarios.js"), ("js", "app.js")):
            source = _asset(*asset)
            visible = re.sub(r"//[^\n]*", "", source)
            visible = re.sub(r"/\*.*?\*/", "", visible, flags=re.S)
            assert "deterministic template" not in visible, asset
